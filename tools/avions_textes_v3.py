#!/usr/bin/env python3
"""Étape V3 (une fois) : remplit les colonnes jaunes du classeur avec la rédaction de Claude.

N'écrit une cellule que si elle est VIDE : les saisies de Florian ne sont jamais écrasées.
Ensuite, le classeur est la référence (avions_classeur.py conserve les colonnes jaunes).
Usage : python tools/avions_textes_v3.py [--classeur data/Avions_wiki.xlsx]
"""
import argparse
import os
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import avions_commun as ac  # noqa: E402

# Noms d'emplacements (glossaire 7.6 étendu). Courts : la colonne fait 190 px sur la fiche.
EMPLACEMENTS_FR = {
    "-": "Crochet d'appontage",
    "Bomb Bay (2x)": "Soute à bombes (×2)",
    "Cargo Bay": "Soute cargo",
    "Cargo Bay (Front)": "Soute cargo (avant)",
    "Cargo Bay (Rear)": "Soute cargo (arrière)",
    "Cargo Bay Front": "Soute cargo (avant)",
    "Cargo Bay Rear": "Soute cargo (arrière)",
    "Cargo Bay(Rear)": "Soute cargo (arrière)",
    "Center Pylon": "Pylône central",
    "CenterLine Pylon": "Pylône axial",
    "Centreline Pylon": "Pylône axial",
    "Chin Turret": "Tourelle de nez",
    "Combined Weapon Bay": "Soute combinée",
    "Combined Weapon Bays": "Soutes combinées",
    "Door guns": "Mitrailleuses de porte",
    "Dorsal Mount": "Support dorsal",
    "Floor Turret Mount": "Tourelle ventrale",
    "Forward Weapon Bay": "Soute avant",
    "Front Fuselage Pylons": "Pylônes de fuselage avant",
    "Front Turret": "Tourelle avant",
    "Fuselage Pylons": "Pylônes de fuselage",
    "Gun Mount": "Support de canon",
    "Heater Bays": "Soutes à missiles IR",
    "Inner Fuselage Pylons": "Pylônes de fuselage int.",
    "Inner L Bay": "Soute intérieure gauche",
    "Inner Left Pylon": "Pylône intérieur gauche",
    "Inner R Bay": "Soute intérieure droite",
    "Inner Right Pylon": "Pylône intérieur droit",
    "Inner Weapon Bays (Combined)": "Soutes int. (combinées)",
    "Inner Weapon Bays (Front)": "Soutes int. (avant)",
    "Inner Weapon Bays (Rear)": "Soutes int. (arrière)",
    "Inner Wing Pylons": "Pylônes d'aile intérieurs",
    "Inner wing pylons": "Pylônes d'aile intérieurs",
    "Internal 35mm Autocannons": "Canons internes de 35 mm",
    "Internal Bays": "Soutes internes",
    "Internal Cannon": "Canon interne",
    "Internal Gun": "Canon interne",
    "Internal Guns": "Canons internes",
    "Internal Laser Mount": "Laser interne",
    "Internal Weapon Bays": "Soutes internes",
    "Internal guns": "Canons internes",
    "JammingPod1": "Nacelle de brouillage",
    "Left Bay": "Soute gauche",
    "Left Fuselage Pylon": "Pylône de fuselage gauche",
    "Left Internal Bay": "Soute interne gauche",
    "Left Internal Cannons": "Canons internes gauches",
    "Left Sponson Pylon": "Pylône de flanc gauche",
    "Left Stub Tip Pylon": "Pylône d'extrémité gauche",
    "Left Weapon Bay": "Soute gauche",
    "Mid Wing Pylons": "Pylônes d'aile médians",
    "Middle Wing Pylons": "Pylônes d'aile médians",
    "Mission Bay": "Soute de mission",
    "Outer Fuselage Pylons": "Pylônes de fuselage ext.",
    "Outer L Bay": "Soute extérieure gauche",
    "Outer Left Pylon": "Pylône extérieur gauche",
    "Outer Left Wing Pylon": "Pylône d'aile ext. gauche",
    "Outer R Bay": "Soute extérieure droite",
    "Outer Right Pylon": "Pylône extérieur droit",
    "Outer Right Wing Pylon": "Pylône d'aile ext. droit",
    "Outer Wing Pylons": "Pylônes d'aile extérieurs",
    "Outer weapon bays": "Soutes extérieures",
    "Outer wing pylons": "Pylônes d'aile extérieurs",
    "Radome": "Radôme",
    "Rear Fuselage Pylons": "Pylônes de fuselage arrière",
    "Rear Turret": "Tourelle arrière",
    "Rear Weapon Bay": "Soute arrière",
    "Right Bay": "Soute droite",
    "Right Fuselage Pylon": "Pylône de fuselage droit",
    "Right Internal Bay": "Soute interne droite",
    "Right Internal Cannon": "Canon interne droit",
    "Right Sponson Pylon": "Pylône de flanc droit",
    "Right Stub Tip Pylon": "Pylône d'extrémité droit",
    "Right Weapon Bay": "Soute droite",
    "Side Weapon Bays": "Soutes latérales",
    "Slingload Hook": "Crochet d'élingue",
    "Stub Pylons": "Pylônes d'ailette",
    "TGP": "Nacelle de désignation",
    "Tail Hook": "Crochet d'appontage",
    "Tailhook": "Crochet d'appontage",
    "Ventral Hardpoint": "Point d'emport ventral",
    "Weapon Bay": "Soute d'armement",
    "Weapon Bays (inner)": "Soutes (intérieures)",
    "Weapon Bays (outer)": "Soutes (extérieures)",
    "Wing Glove Pylons": "Pylônes d'emplanture",
    "Wing Pylons": "Pylônes d'aile",
    "Wingtip": "Saumons d'aile",
    "Wingtip Pylons": "Saumons d'aile",
}

# Rôles (vocabulaire fermé 7.4, 3 au plus) : tirés de la description du jeu ; vide si elle ne permet pas de trancher.
ROLES = {
    "COIN": ["Contre-insurrection", "Appui aérien rapproché"],
    "AttackHelo1": ["Hélicoptère d'attaque"],
    "trainer": ["Entraînement", "Attaque au sol"],
    "UtilityHelo1": ["Hélicoptère utilitaire", "Transport"],
    "VTOLTrainer1": ["Entraînement", "Attaque au sol", "Embarqué (porte-avions)"],
    "CAS1": ["Appui aérien rapproché", "Attaque au sol", "Lutte antinavire"],
    "QuadVTOL1": ["Transport"],
    "Fighter1": ["Supériorité aérienne"],
    "SmallFighter1": ["Multirôle", "Embarqué (porte-avions)"],
    "EW1": ["Guerre électronique", "SEAD"],
    "Multirole1": ["Multirôle", "Embarqué (porte-avions)"],
    "Darkreach": ["Bombardier furtif"],
    "FastBomber1": ["Bombardier", "Guerre électronique"],
    "Aryx_MiG-15": [],
    "Aryx_PropAttacker1": ["Contre-insurrection", "Attaque au sol", "Entraînement"],
    "Aryx_LightHelicopter1": ["Reconnaissance", "Hélicoptère d'attaque"],
    "Aryx_LightFighter1": ["Supériorité aérienne", "Chasseur"],
    "Aryx_F16M_KingViper": [],
    "1509_palafighter1": ["Chasseur"],
    "blacknight2u_F117A_Nighthawk": ["Frappe furtive"],
    "P_Trisurface1": ["Multirôle"],
    "Aryx_CargoPlane1": ["Transport", "Bombardier"],
    "Aryx_F22E_StrikeRaptor": ["Supériorité aérienne", "Attaque au sol"],
    "Aryx_Interceptor1": ["Intercepteur", "Embarqué (porte-avions)"],
}

# Descriptions : traduction fidèle et complète de la description du jeu.
DESCRIPTIONS = {
    "COIN": "Conçu pour la contre-insurrection et les conflits de faible intensité, le CI-22 Cricket dispose d'un train "
            "d'atterrissage robuste et de capacités de décollage et d'atterrissage courts, qui lui permettent d'opérer depuis "
            "des terrains non préparés. S'il manque de vitesse et de capacité de survie, le CI-22 offre une agilité remarquable "
            "et peut emporter un large éventail d'armes.",
    "AttackHelo1": "Le SAH-46 Chicane conserve la configuration classique d'un hélicoptère d'attaque tout en intégrant des "
                   "améliorations qui en font toujours une présence redoutable sur le champ de bataille de la fin du XXIe siècle. "
                   "Sa technologie furtive de pointe lui permet de lancer des attaques surprises depuis l'abri du relief, puis de "
                   "disparaître des radars sans laisser de trace.",
    "trainer": "Avion d'entraînement avancé à réaction, bien adapté aux missions de combat léger, le T/A-30 est très manœuvrant "
               "à basse et moyenne vitesse et pardonne les erreurs de pilotage. Particulièrement à l'aise dans les frappes à basse "
               "altitude, le Compass dispose de gouvernes redondantes et de deux réacteurs qui améliorent sa capacité de survie.",
    "UtilityHelo1": "Conçu pour un programme PALA de transport vertical rapide dans les années 2050, l'UH-90 a été pensé comme un "
                    "hélicoptère utilitaire combiné de taille moyenne, mais il a depuis reçu de nombreux points d'emport qui lui "
                    "permettent de servir d'appareil d'appui feu. Doté d'un rotor rigide coaxial et de deux soufflantes carénées "
                    "turbo-électriques réversibles en configuration propulsive, l'Ibis atteint rapidement de hautes vitesses de "
                    "croisière avant de décélérer jusqu'à un atterrissage vertical en moins de 30 secondes.",
    "VTOLTrainer1": "Doté d'un système unique de poussée canalisée à trois tuyères, cet avion d'entraînement sert aussi d'avion "
                    "d'attaque économique à décollage vertical. Le Vagrant est très manœuvrant et atteint des vitesses transsoniques ; "
                    "il dispose aussi d'un radar qui lui permet d'engager des cibles aériennes au-delà de la portée visuelle. Bien "
                    "qu'il n'ait pas de furtivité avancée, le VT-7 reste la colonne vertébrale de l'aéronavale de la BDF.",
    "CAS1": "Le Brawler est une plateforme fiable et robuste, conçue pour emporter de grandes quantités d'armes de précision. Ses "
            "propfans de pointe allient endurance et vitesse respectable, tandis que ses gouvernes redondantes, ses moteurs isolés "
            "et le lourd blindage de son cockpit lui assurent une excellente capacité de survie. L'A-19 est dans son élément face "
            "aux convois de véhicules, aux défenses antiaériennes courte portée dispersées ou aux navires opérant en eaux côtières.",
    "QuadVTOL1": "Alliant capacité de levage vertical et 20 tonnes de capacité d'emport, ce grand quadrimoteur à ailes basculantes "
                 "peut livrer des approvisionnements vitaux aux positions avancées ou servir d'appareil d'appui feu polyvalent. Ses "
                 "4 rotors sont reliés par des arbres de transmission internes, ce qui permet au Tarantula de continuer à voler "
                 "après avoir encaissé des dégâts importants.",
    "Fighter1": "Le FS-12 Revoker est conçu avant tout pour l'agilité : il est capable de manœuvres extrêmes à presque toutes les "
                "vitesses. Ses soutes internes réduisent la traînée et lui permettent d'atteindre des vitesses supersoniques même à "
                "basse altitude et lourdement armé. Moins furtif que les chasseurs plus récents, le Revoker reste un moyen économique "
                "d'obtenir la supériorité aérienne.",
    "SmallFighter1": "Chasseur multirôle compact et avancé, le Vortex peut opérer à la verticale depuis de petits porte-avions et "
                     "rester indétecté grâce à une furtivité exceptionnelle. Rapide, agile et furtif quand il est peu chargé, il "
                     "perd une partie de ces qualités sous une lourde charge d'armement.",
    "EW1": "L'EW-25 est un avion de soutien à décollage court et atterrissage vertical (STOVL) doté de puissants systèmes de guerre "
           "électronique. Capable de trianguler la position des radars, de brouiller les installations radar ennemies, de neutraliser "
           "les défenses antiaériennes et même d'abattre des missiles avec un laser à haute énergie, un seul Medusa peut facilement "
           "renverser le cours d'une bataille.",
    "Multirole1": "Grand chasseur multirôle biréacteur apte à l'embarquement, l'Ifrit convient aussi bien à la supériorité aérienne "
                  "qu'à l'attaque au sol. Il offre une poussée et une maniabilité excellentes, une signature radar réduite et des "
                  "capteurs puissants. Deux soutes internes, plus deux soutes latérales pour missiles à guidage infrarouge, permettent "
                  "à l'Ifrit d'emporter une charge interne importante sans nuire à sa furtivité ni à sa vitesse de pointe.",
    "Darkreach": "Avec ses moteurs puissants, sa silhouette furtive et ses 4 grandes soutes internes, le SFB-81 Darkreach est un "
                 "bombardier redoutable capable de délivrer des armes conventionnelles comme nucléaires. Son aile volante intégrée lui "
                 "permet d'opérer dans des régimes de vol variés, de la pénétration à très basse altitude au bombardement stratégique "
                 "à distance de sécurité.",
    "FastBomber1": "L'Alkyon AB-4 est un bombardier supersonique avancé doté de capacités de guerre électronique, conçu pour délivrer "
                   "discrètement des charges nucléaires ou conventionnelles sur des cibles de grande valeur. Ses 4 réacteurs à "
                   "postcombustion et ses ailes à géométrie variable lui donnent une vitesse de pointe incroyable et assez de "
                   "maniabilité pour échapper aux menaces aériennes. En cas de dégâts critiques, une capsule d'évacuation permet à "
                   "l'équipage de s'éjecter en sécurité à toute altitude et à toute vitesse.",
    "Aryx_MiG-15": "Des MiG-15 en plein '85 :wilt: :sob:",
    "Aryx_PropAttacker1": "Le Hartnell-Wesley OA-27 Cavalier est un avion léger d'entraînement, d'attaque et de contre-insurrection "
                          "en service dans l'État primévien et la PALAAF. Conçu avant la révolution dans l'un des micro-États "
                          "primévien, le Cavalier est l'appareil de choix pour les opérations d'attaque de faible intensité, les "
                          "escarmouches frontalières et la formation intermédiaire des pilotes. Il est fiable, robuste et "
                          "remarquablement peu coûteux à acquérir et à entretenir.",
    "Aryx_LightHelicopter1": "Le RAH-72 est un hélicoptère de reconnaissance et d'attaque légère. Capable d'emporter un bon nombre de "
                             "munitions légères, c'est un monoplace biturbine à rotors coaxiaux et empennage en V, avec des soutes "
                             "internes sous le fuselage et quatre pylônes externes sur des ailettes dont les extrémités portent ses "
                             "lance-leurres. Il bénéficie d'une mise en forme à facettes peu observable, d'excellents capteurs "
                             "optiques et d'un grand radar de triangulation au sol monté au sommet du mât rotor, qui détecte unités "
                             "terrestres et navires à longue distance.",
    "Aryx_LightFighter1": "Malgré une conception ancienne, ce chasseur léger supersonique est exceptionnellement agile et très adapté "
                          "aux missions de supériorité aérienne, avec des options d'attaque au sol limitées à basse altitude. Ses deux "
                          "réacteurs lui donnent un rapport poussée/poids et une vitesse de pointe incroyables, tandis qu'une soute "
                          "interne axiale réduit la traînée de la charge.",
    "Aryx_F16M_KingViper": "Bien que plus que centenaire, le F-16 reste un pilier des forces aériennes du monde entier, ses dernières "
                           "variantes Block 165 étant en service aussi bien dans la BDAF que dans la PALAAF. Les programmes de "
                           "modernisation et un lobbying acharné l'ont maintenu dans la course : le Viper est vraiment éternel.",
    "1509_palafighter1": "Une conception plus ancienne mais éprouvée, qui remonte à la révolution primévienne. L'Agni est un chasseur "
                         "monoréacteur tout temps doté d'un grand nombre de points d'emport pour plus de souplesse. Sans aucune soute "
                         "interne, l'appareil est plus facile à détecter que les modèles plus modernes.",
    "blacknight2u_F117A_Nighthawk": "Avion de frappe subsonique rapide à très faible observabilité, doté de deux réacteurs F404 sans "
                                    "postcombustion, de commandes de vol électriques numériques, d'un système de désignation "
                                    "électro-optique et de deux soutes internes. Un large choix de chargements de bombes et de missiles "
                                    "compatibles offre un jeu varié. Soutes fermées, les armes internes restent masquées au radar ; "
                                    "ouvrir une soute ou sortir le train augmente la signature radar. Un brouilleur radar actif fixe est "
                                    "installé dans chaque chargement : sélectionne-le comme arme, désigne une cible radar suivie et "
                                    "maintiens le tir pour brouiller. Maintenir le frein de roues pendant un roulage d'atterrissage "
                                    "valide déploie le parachute de freinage (à usage unique) une fois les trois roues au sol et "
                                    "l'appareil aligné sur sa trajectoire.",
    "P_Trisurface1": "Conçu au tournant du siècle, le FS-3 Ternion est un grand chasseur multirôle biréacteur. D'abord développé pour "
                     "de longues patrouilles au-dessus du Pacifique, il a ensuite été adapté à des besoins opérationnels croissants. "
                     "Malgré l'absence de furtivité, des modernisations régulières ont permis au Ternion de rester un défenseur "
                     "fiable du Heartland.",
    "Aryx_CargoPlane1": "Plus gros appareil d'Aryx Dynamics à ce jour, le MC-260 Chimera est un avion cargo lourd polyvalent, conçu pour "
                        "des opérations multirôles lourdes et rustiques dans des conditions spartiates, près des lignes de front. Il "
                        "offre des performances de décollage et d'atterrissage courts à faible charge grâce à son aile soufflée et à "
                        "ses inverseurs de poussée, et tient dans les hangars moyens existants grâce à son gouvernail à volets "
                        "divisés et à ses dérives ventrales uniques. Le Chimera peut servir de canonnière, de bombardier, d'avion de "
                        "guet aérien (AEW) et de transport de fret.",
    "Aryx_F22E_StrikeRaptor": "Proposé à la fin des années 2040, en pleine Troisième Guerre mondiale, comme programme de modernisation "
                              "et de refonte du F-22C d'exportation, le F-22E Strike Raptor reprend la forme bien connue du F-22 en y "
                              "ajoutant une paire d'EOTS intégrés aux arêtes du fuselage central pour guider les munitions de "
                              "précision, ainsi qu'une soute agrandie et approfondie pour emporter de gros volumes d'armement. "
                              "Excellente agilité, poussée vectorielle, nacelles d'armes furtives possibles sur ses pylônes d'aile "
                              "intérieurs, et désormais tout à fait capable d'attaque au sol comme de domination aérienne : le F-22E "
                              "est aussi féroce que l'était son modèle original un demi-siècle plus tôt.",
    "Aryx_Interceptor1": "Le FS-41 Eclipse est un grand intercepteur embarqué biréacteur supersonique, conçu pour localiser, intercepter "
                         "et détruire les cibles de grande valeur avant qu'elles ne menacent l'espace aérien allié et les groupes "
                         "aéronavals. Il allie des vitesses supérieures à Mach 2, un radar extraordinairement puissant et une furtivité "
                         "impressionnante pour son rôle d'interception et de domination, renforcés par son armement emblématique, "
                         "le missile AAM-45 Sabre. Ses ailes à géométrie variable le rendent en outre très agile, et il peut emporter "
                         "une charge air-sol limitée pour des missions variées.",
}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--classeur", default=ac.CLASSEUR)
    args = ap.parse_args()
    for cle, roles in ROLES.items():
        assert all(r in ac.ROLES for r in roles) and len(roles) <= 3, cle
    wb = openpyxl.load_workbook(args.classeur)
    ws = wb["Avions"]
    ent = [c.value for c in ws[1]]
    ecrits = 0
    for l in ws.iter_rows(min_row=2):
        cle = l[0].value
        for col, val in (("Rôle(s)", " · ".join(ROLES.get(cle, [])) or None), ("Description (FR)", DESCRIPTIONS.get(cle))):
            c = l[ent.index(col)]
            if val and not ac.texte(c.value):
                c.value = val
                ecrits += 1
    ws = wb["Emports"]
    ent = [c.value for c in ws[1]]
    manquants = set()
    for l in ws.iter_rows(min_row=2):
        nom = l[ent.index("Nom (jeu)")].value
        c = l[ent.index("Nom (FR)")]
        if not ac.texte(c.value):
            if nom in EMPLACEMENTS_FR:
                c.value = EMPLACEMENTS_FR[nom]
                ecrits += 1
            else:
                manquants.add(nom)
    wb.save(args.classeur)
    print(f"{ecrits} cellule(s) remplie(s).")
    if manquants:
        print("Noms d'emplacement non traduits :", ", ".join(sorted(manquants)))
    vides = [c for c, r in ROLES.items() if not r]
    print("Rôles laissés vides :", ", ".join(vides) or "aucun")


if __name__ == "__main__":
    main()
