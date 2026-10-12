#!/usr/bin/env python3
"""Tests des scripts du wiki avions (critères V2 et V4 du plan). Usage : python tools/test_avions.py"""
import json
import os
import shutil
import subprocess
import sys
import tempfile

import openpyxl

OUTILS = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, OUTILS)
import avions_commun as ac  # noqa: E402

echecs = 0
verifs = 0


def ok(cond, quoi):
    global echecs, verifs
    verifs += 1
    if not cond:
        echecs += 1
        print("  ÉCHEC :", quoi)


def contenu(chemin):
    wb = openpyxl.load_workbook(chemin)
    return {ws.title: [[c.value for c in ligne] for ligne in ws.iter_rows()] for ws in wb.worksheets}


def lancer(script, *args):
    r = subprocess.run([sys.executable, os.path.join(OUTILS, script), *args], capture_output=True, text=True)
    if r.returncode != 0:
        print(r.stdout, r.stderr)
        raise SystemExit(f"{script} a échoué")
    return r.stdout


def ligne_avion(chemin, cle):
    ws = openpyxl.load_workbook(chemin)["Avions"]
    ent = [c.value for c in ws[1]]
    for l in ws.iter_rows(min_row=2, values_only=True):
        if l[0] == cle:
            return dict(zip(ent, l))
    return None


def main():
    tmp = tempfile.mkdtemp()
    try:
        x = os.path.join(tmp, "Avions_wiki.xlsx")
        j = os.path.join(tmp, "aircraft.json")
        # Base : classeur du dépôt s'il existe (pour garder ses textes), sinon vierge.
        if os.path.exists(ac.CLASSEUR):
            shutil.copy2(ac.CLASSEUR, x)
        lancer("avions_classeur.py", "--sortie", x)
        c1 = contenu(x)
        lancer("avions_classeur.py", "--sortie", x)
        ok(contenu(x) == c1, "V2 : deux exécutions donnent le même contenu")

        avions = c1["Avions"]
        ok(len(avions) - 1 == 24, f"V2 : 24 appareils (obtenu {len(avions) - 1})")
        ok(all(l[0] for l in avions[1:]), "V2 : chaque appareil a une clé")
        ok("UFO" not in [l[0] for l in avions[1:]], "A6 : OVNI exclu")

        # Saisies de Florian : texte perso, nom FR, corrections (une valide, une à champ inconnu, une à clé inconnue).
        wb = openpyxl.load_workbook(x)
        ws = wb["Avions"]
        ent = [c.value for c in ws[1]]
        for l in ws.iter_rows(min_row=2):
            if l[0].value == "COIN":
                l[ent.index("Texte perso")].value = "Mon avion préféré"
        we = wb["Emports"]
        ente = [c.value for c in we[1]]
        for l in we.iter_rows(min_row=2):
            if l[0].value == "COIN" and l[ente.index("Nom (jeu)")].value == "Inner Wing Pylons":
                l[ente.index("Nom (FR)")].value = "Test pylônes"
        wc = wb["Corrections"]
        wc.append(["COIN", None, "Vitesse max. (km/h)", 3000, "test"])
        wc.append(["COIN", None, "Champ bidon", 1, None])
        wc.append(["XXX", None, "Rang requis", 1, None])
        wc.append(["COIN", "Inner Wing Pylons", "Charge max / point (kg)", 250, None])
        wb.save(x)

        lancer("avions_classeur.py", "--sortie", x)
        ci = ligne_avion(x, "COIN")
        ok(ci["Texte perso"] == "Mon avion préféré", "V2 : Texte perso conservé")
        c2 = contenu(x)
        ok(any(l[0] == "COIN" and l[4] == "Test pylônes" for l in c2["Emports"]), "V2 : Nom (FR) conservé")
        corr = c2["Corrections"][1:]
        ok(len(corr) == 4 and corr[0][:4] == ["COIN", None, "Vitesse max. (km/h)", 3000], "V2 : lignes de Corrections conservées")
        ok(corr[0][5] == "appliquée", f"A20 : correction valide appliquée ({corr[0][5]})")
        ok(str(corr[1][5]).startswith("ignorée"), f"A20 : champ inconnu ignoré ({corr[1][5]})")
        ok(str(corr[2][5]).startswith("ignorée"), f"A20 : clé inconnue ignorée ({corr[2][5]})")
        ok(corr[3][5] == "appliquée", f"A20 : correction d'emplacement appliquée ({corr[3][5]})")
        ok(ci["Vitesse max. (km/h)"] == 3000 and ci["Note vitesse"] == 10.0, "V4 : vitesse forcée → note 10,0")
        ok(ligne_avion(x, "Aryx_Interceptor1")["Note vitesse"] <= 9.6, "V4 : FS-41 ≤ 9,6 après correction du CI-22")

        # JSON (V4), si le script existe.
        if os.path.exists(os.path.join(OUTILS, "avions_vers_json.py")):
            lancer("avions_vers_json.py", "--classeur", x, "--sortie", j)
            d1 = open(j, encoding="utf-8").read()
            lancer("avions_vers_json.py", "--classeur", x, "--sortie", j)
            d2 = open(j, encoding="utf-8").read()
            sans_date = lambda s: "\n".join(l for l in s.splitlines() if '"genere_le"' not in l)  # noqa: E731
            ok(sans_date(d1) == sans_date(d2), "V4 : sortie identique sur deux exécutions")
            data = json.loads(d2)
            a = data["appareils"]
            ok(a["COIN"]["vitesse_kmh"] == 3000 and "vitesse_kmh" in a["COIN"]["corriges"], "V4 : correction dans le JSON")
            ok(a["COIN"]["note_vitesse"] == 10.0, "V4 : note vitesse CI-22 = 10,0")
            ok(a["Aryx_Interceptor1"]["note_vitesse"] <= 9.6, "V4 : note vitesse FS-41 ≤ 9,6")
            ok(a["COIN"]["emports_forces"]["Inner Wing Pylons"]["charge_kg"] == 250, "V4 : emport forcé")
            ok(a["COIN"]["texte_perso"] == "Mon avion préféré", "V4 : texte perso")
            ok(data["emplacements_fr"]["COIN"]["Inner Wing Pylons"] == "Test pylônes", "V4 : nom FR d'emplacement")
            ok(data["alias"].get("CI-22 Cricket") == "COIN", "V4 : alias unitName → clé")

        # Contrôles chiffrés du plan (V4), sur le classeur sans les corrections de test.
        m = ac.construire(ac.lire_classeur(ac.CLASSEUR) if os.path.exists(ac.CLASSEUR) else None, journal=lambda *_: None)
        p = {a["cle"]: a for a in m["appareils"]}
        ok(p["COIN"]["taille"] == "S" and p["COIN"]["note_rcs"] == 4 and p["COIN"]["maniabilite"] == 6.5, "V4 : CI-22 S, RCS 4, agilité 6,5")
        ok(p["Aryx_CargoPlane1"]["taille"] == "H", "V4 : MC-260 taille H")
        ok(p["blacknight2u_F117A_Nighthawk"]["note_rcs"] == 0, "V4 : F-117A RCS 0")
        ok(p["Aryx_Interceptor1"]["note_vitesse"] == 10.0, "V4 : FS-41 vitesse 10,0")
        ok(p["COIN"]["laser_cibles"] == 2 and p["Aryx_MiG-15"]["laser_cibles"] == 0, "A21 : désignateur CI-22 = 2, MiG-15 = 0")
        tailles = [a["taille"] for a in m["appareils"]]
        ok((tailles.count("S"), tailles.count("M"), tailles.count("H")) == (12, 8, 4), f"A7 : 12 S, 8 M, 4 H ({tailles.count('S')}, {tailles.count('M')}, {tailles.count('H')})")
        ok(all(not s["affiche"] for s in p["Aryx_CargoPlane1"]["emplacements"] if s["nom_jeu"] == "-"), "A23 : « - » du MC-260 masqué")
        ok(ac.note_rcs(5e-7) == 0 and ac.note_rcs(0.001) == 1 and ac.note_rcs(0.02) == 2 and ac.note_rcs(0.08) == 3
           and ac.note_rcs(0.5) == 4 and ac.note_rcs(5) == 5, "A12 : échelle RCS")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    print(f"OK ({verifs} vérifications)" if echecs == 0 else f"ÉCHEC : {echecs} / {verifs}")
    return 1 if echecs else 0


if __name__ == "__main__":
    sys.exit(main())
