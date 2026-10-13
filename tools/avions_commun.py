"""Code commun aux scripts du wiki avions (plan wiki-avions-no, v3.2).

construire() fusionne NO-stats.xlsx et le dernier export du mod, applique la feuille « Corrections » (A20),
puis calcule tailles, classes de charge, notes et seuils (A7, A9, A12, A14, A15).
Utilisé par avions_classeur.py (classeur Avions_wiki.xlsx) et avions_vers_json.py (aircraft.json).
"""
import csv
import glob
import math
import os
import re
from decimal import ROUND_HALF_UP, Decimal

RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
NO_STATS = os.path.join(RACINE, "data", "sources", "NO-stats.xlsx")
CLASSEUR = os.path.join(RACINE, "data", "Avions_wiki.xlsx")
DUMP = os.path.join(RACINE, "data", "dump")

# ---------------------------------------------------------------------------------------------------------------
# Colonnes attendues dans NO-stats (feuille « Avions ») : cherchées par en-tête, arrêt clair si l'une manque.
COLONNES_NO = {
    "source": "Source", "nom": "Appareil", "code": "Code", "type": "Type", "rang": "Rang requis", "cout_m": "Coût (M$)",
    "poussee_kn": "Poussée totale des réacteurs (kN)", "puissance_kw": "Puissance totale des turbomoteurs (kW)",
    "masse_vide_kg": "Masse à vide (kg)", "masse_max_kg": "Masse max. (kg)", "carburant_kg": "Carburant interne (kg)",
    "vitesse_kmh": "Vitesse max. (km/h)", "decrochage_kmh": "Vitesse de décrochage (km/h)",
    "maniabilite": "Maniabilité (0 à 10)", "g_max": "Facteur G max", "decollage_vertical": "Décollage vertical",
    "decollage_ms": "Vitesse de décollage (m/s)", "decollage_m": "Distance de décollage (m)",
    "dimensions": "Dimensions L × l × h (m)", "canons_fixes": "Canons fixes (nombre)",
    "radar_m": "Portée radar embarqué (m)", "brouilleur": "Brouilleur radar", "leurres": "Leurres IR (nombre)",
    "paillettes": "Paillettes (nombre)", "blindage": "Palier de blindage", "rcs_m2": "Signature radar (valeur jeu)",
    "description_jeu": "Description (jeu)", "notes": "Notes",
    "role_surface": "Rôle : anti-surface (0 à 1)", "role_air": "Rôle : anti-air (0 à 1)",
    "role_missile": "Rôle : anti-missile (0 à 1)", "role_radar": "Rôle : anti-radar (0 à 1)",
}

# Vocabulaire fermé des rôles (plan 7.4).
ROLES = [
    "Chasseur", "Intercepteur", "Supériorité aérienne", "Multirôle", "Attaque au sol", "Appui aérien rapproché",
    "Contre-insurrection", "Bombardier", "Bombardier furtif", "Frappe furtive", "Guerre électronique", "SEAD",
    "Hélicoptère d'attaque", "Hélicoptère utilitaire", "Transport", "Entraînement", "Reconnaissance",
    "Lutte antinavire", "Embarqué (porte-avions)",
]

CATEGORIES = ["CANON", "MISSILE AA", "MISSILE AS", "ROQUETTE", "BOMBE", "NUC", "LASER", "BROUILLEUR", "CARGO"]
# Pastilles de la fiche (amendement v3.3, choix de Florian ; EW pour le brouilleur, non précisé).
ACRONYMES = {"CANON": "CN", "MISSILE AA": "AAM", "MISSILE AS": "ASM", "ROQUETTE": "RK", "BOMBE": "BMB", "NUC": "NK",
             "LASER": "LZ", "BROUILLEUR": "EW", "CARGO": "TR"}

# Champs corrigeables (plan 7.7, amendé v3.2) : libellé → (champ JSON, type). N nombre, T texte, B OUI/NON, L catégories.
CHAMPS_APPAREIL = [
    ("Coût (M$)", "cout_m", "N"), ("Rang requis", "rang", "N"), ("Vitesse max. (km/h)", "vitesse_kmh", "N"),
    ("Décrochage (km/h)", "decrochage_kmh", "N"), ("Maniabilité (0 à 10)", "maniabilite", "N"),
    ("Facteur G max", "g_max", "N"), ("Masse à vide (kg)", "masse_vide_kg", "N"), ("Masse max. (kg)", "masse_max_kg", "N"),
    ("Charge utile (kg)", "charge_utile_kg", "N"), ("Carburant (kg)", "carburant_kg", "N"), ("Poussée (kN)", "poussee_kn", "N"),
    ("Puissance (kW)", "puissance_kw", "N"), ("Décollage vertical", "decollage_vertical", "B"),
    ("Distance de décollage (m)", "decollage_m", "N"), ("Vitesse de décollage (m/s)", "decollage_ms", "N"),
    ("Portée radar (m)", "radar_m", "N"), ("Brouilleur", "brouilleur", "B"), ("Canons fixes", "canons_fixes", "N"),
    ("Leurres IR", "leurres", "N"), ("Paillettes", "paillettes", "N"), ("Palier de blindage", "blindage", "N"),
    ("Signature radar (m²)", "rcs_m2", "N"), ("Désignateur laser (cibles)", "laser_cibles", "N"),
    ("IR à 30 % de gaz", "ir_force.ir_30", "N"), ("IR plein gaz", "ir_force.ir_max", "N"), ("Taille (S/M/H)", "taille", "T"),
    ("Type", "type", "T"), ("Longueur (m)", "longueur_m", "N"), ("Envergure (m)", "envergure_m", "N"),
]
CHAMPS_EMPLACEMENT = [
    ("Points", "points", "N"), ("Soute", "soute", "B"), ("Charge max / point (kg)", "charge_kg", "N"),
    ("Catégories", "categories", "L"), ("Nb armes compatibles", "nb_armes", "N"), ("Masquer l'emplacement", "masquer", "B"),
]
CHAMPS = {lib: (champ, typ, False) for lib, champ, typ in CHAMPS_APPAREIL}
CHAMPS.update({lib: (champ, typ, True) for lib, champ, typ in CHAMPS_EMPLACEMENT})


# ---------------------------------------------------------------------------------------------------------------
def arrondi(v, decimales=1):
    """Arrondi « au plus proche, 0,5 vers le haut » (et non l'arrondi bancaire de Python)."""
    if v is None:
        return None
    q = Decimal(1).scaleb(-decimales)
    return float(Decimal(repr(v)).quantize(q, rounding=ROUND_HALF_UP))


def nombre(v):
    """Cellule ou champ CSV → float ; vide, « N/A », texte → None."""
    if v is None:
        return None
    if isinstance(v, bool):
        return float(v)
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "").replace(" ", "").replace(",", ".")
    if s == "":
        return None
    try:
        return float(s)
    except ValueError:
        return None


def booleen(v):
    if v is None:
        return None
    if isinstance(v, bool):
        return v
    s = str(v).strip().lower()
    if s in ("oui", "o", "vrai", "true", "1", "yes", "x"):
        return True
    if s in ("non", "n", "faux", "false", "0", "no"):
        return False
    return None


def texte(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def entier_si_rond(v):
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


# ---------------------------------------------------------------------------------------------------------------
def lire_no_stats(chemin=NO_STATS):
    import openpyxl
    wb = openpyxl.load_workbook(chemin, data_only=True, read_only=True)
    if "Avions" not in wb.sheetnames:
        raise SystemExit(f"{chemin} : feuille « Avions » introuvable.")
    lignes = list(wb["Avions"].iter_rows(values_only=True))
    entetes = [texte(h) for h in lignes[0]]
    index = {}
    for cle, entete in COLONNES_NO.items():
        if entete not in entetes:
            raise SystemExit(f"{chemin} : colonne « {entete} » introuvable dans la feuille Avions "
                             f"(NO-stats a peut-être été modifié : renomme la colonne ou adapte COLONNES_NO).")
        index[cle] = entetes.index(entete)
    appareils = []
    for ligne in lignes[1:]:
        if not ligne or not texte(ligne[index["nom"]]):
            continue
        appareils.append({cle: ligne[i] for cle, i in index.items()})
    return appareils


def dernier_export(motif):
    fichiers = sorted(glob.glob(os.path.join(DUMP, motif)))
    if not fichiers:
        raise SystemExit(f"Aucun export {motif} dans data/dump/ : lance l'export des appareils en jeu (F1 → Diagnostic).")
    return fichiers[-1]


def lire_csv(chemin):
    with open(chemin, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f, delimiter=";"))


def est_aeronef(ligne_export):
    """A22 : un aéronef a une voilure tournante, une hélice ou une tuyère, et ni propulsion navale ni roues motrices."""
    moteurs = set((ligne_export.get("moteurs") or "").split("|"))
    if moteurs & {"PlayerShipPropulsion", "WheelEngine"}:
        return False
    return bool(moteurs & {"JetNozzle", "ConstantSpeedProp", "PropFan", "RotorShaft"})


def est_exclu(ligne_export):
    """A6 (OVNI, contenu événementiel) et A22 (navires, véhicules)."""
    return (ligne_export.get("evenementiel") == "1" or ligne_export.get("json_key") == "UFO"
            or not est_aeronef(ligne_export))


def type_depuis_moteurs(moteurs):
    m = set((moteurs or "").split("|"))
    if "RotorShaft" in m:
        return "Hélicoptère"
    if "JetNozzle" in m:
        return "Avion à réaction"
    return "Avion à hélice"


TYPES_COURTS = [("avion à hélice", "HÉLICE"), ("hélicoptère", "HÉLICO"), ("avion à réaction", "RÉACTION"),
                ("avion à décollage vertical", "VTOL"), ("convertible à rotors basculants", "CONVERTIBLE")]


def type_court(t):
    """Plan 7.8 : puce de type."""
    if not t:
        return None
    bas = t.strip().lower()
    if "géométrie variable" in bas:
        return "GÉOM. VARIABLE"
    for debut, court in TYPES_COURTS:
        if bas.startswith(debut):
            return court
    return t.strip().upper()[:14]


def dimensions(v):
    """« 10.4 × 11.6 × 3.4 » → (10.4, 11.6, 3.4)."""
    if v is None:
        return (None, None, None)
    parts = [nombre(p) for p in re.split(r"[×xX*]", str(v))]
    parts += [None] * (3 - len(parts))
    return tuple(parts[:3])


# ---------------------------------------------------------------------------------------------------------------
def tiers_log(valeurs):
    """A7 / A9 : bornes min, max et seuils S/M des tiers égaux en échelle log."""
    vals = [v for v in valeurs if v is not None and v > 0]
    if not vals:
        return None
    vmin, vmax = min(vals), max(vals)
    lmin, lmax = math.log10(vmin), math.log10(vmax)
    return {"min": arrondi(vmin), "max": arrondi(vmax),
            "s_max": arrondi(10 ** (lmin + (lmax - lmin) / 3)), "m_max": arrondi(10 ** (lmin + 2 * (lmax - lmin) / 3))}


def classe(v, seuils):
    """S / M / H selon les seuils (en unités réelles, comme le fait le jeu pour un appareil non référencé)."""
    if v is None or v <= 0 or not seuils:
        return None
    if v <= seuils["s_max"]:
        return "S"
    if v <= seuils["m_max"]:
        return "M"
    return "H"


def note_rang(v, valeurs):
    """A14 : 1 + 9 × (nb ≤ v − 1) / (n − 1), une décimale (ex æquo en haut)."""
    vals = [x for x in valeurs if x is not None]
    if v is None or not vals:
        return None
    n = len(vals)
    if n == 1:
        return 10.0
    nb = sum(1 for x in vals if x <= v)
    return arrondi(1 + 9 * (nb - 1) / (n - 1))


def note_rcs(rcs):
    """A12 : clamp(arrondi(log10(rcs) + 4), 0, 5)."""
    if rcs is None or rcs <= 0:
        return None
    return int(max(0, min(5, arrondi(math.log10(rcs) + 4, 0))))


# ---------------------------------------------------------------------------------------------------------------
def lire_classeur(chemin=CLASSEUR):
    """Colonnes éditables (par clé) et feuille Corrections (ligne à ligne) d'un classeur existant."""
    res = {"avions": {}, "emports": {}, "corrections": [], "stats": {}}
    if not os.path.exists(chemin):
        return res
    import openpyxl
    wb = openpyxl.load_workbook(chemin, data_only=False)
    if "Avions" in wb.sheetnames:
        ws = wb["Avions"]
        ent = [texte(c.value) for c in ws[1]]
        for ligne in ws.iter_rows(min_row=2, values_only=True):
            d = dict(zip(ent, ligne))
            cle = texte(d.get("Clé"))
            if cle:
                res["avions"][cle] = {"roles": texte(d.get("Rôle(s)")), "description_fr": texte(d.get("Description (FR)")),
                                      "texte_perso": texte(d.get("Texte perso")),
                                      "tags": texte(d.get("Tags"))}
                res["stats"][cle] = d
    if "Emports" in wb.sheetnames:
        ws = wb["Emports"]
        ent = [texte(c.value) for c in ws[1]]
        for ligne in ws.iter_rows(min_row=2, values_only=True):
            d = dict(zip(ent, ligne))
            cle, nom = texte(d.get("Clé")), texte(d.get("Nom (jeu)"))
            if cle and nom is not None:
                res["emports"][(cle, nom)] = texte(d.get("Nom (FR)"))
    if "Corrections" in wb.sheetnames:
        ws = wb["Corrections"]
        for ligne in ws.iter_rows(min_row=2, max_col=5, values_only=True):
            ligne = list(ligne) + [None] * (5 - len(ligne))
            if any(v is not None and str(v).strip() != "" for v in ligne):
                res["corrections"].append(ligne[:5])
    return res


# ---------------------------------------------------------------------------------------------------------------
def construire(classeur=None, no_stats=NO_STATS, export_avions=None, export_emports=None, journal=print):
    """Modèle complet : {'appareils': [...], 'seuils': {...}, 'corrections': [...], 'exports': (...)}."""
    classeur = classeur if classeur is not None else lire_classeur()
    export_avions = export_avions or dernier_export("*_avions_v*.csv")
    export_emports = export_emports or dernier_export("*_emports_v*.csv")
    no = lire_no_stats(no_stats)
    exp = {r["json_key"]: r for r in lire_csv(export_avions)}
    par_nom = {r["unit_name"]: r for r in exp.values()}
    lignes_emports = lire_csv(export_emports)

    # --- Appareils référencés : NO-stats appariés par nom (= unit_name), puis aéronefs de l'export absents de NO-stats.
    appareils, vus = [], set()
    for n in no:
        nom = texte(n["nom"])
        e = par_nom.get(nom)
        if e is None:
            journal(f"ERREUR : « {nom} » (NO-stats) absent de l'export du jeu : pas de clé, appareil ignoré.")
            continue
        if est_exclu(e):
            continue
        appareils.append(_appareil(e, n))
        vus.add(e["json_key"])
    for cle, e in sorted(exp.items()):
        if cle not in vus and not est_exclu(e):
            journal(f"Info : {cle} ({e['unit_name']}) absent de NO-stats : valeurs de l'export seulement.")
            appareils.append(_appareil(e, None))
    cles = {a["cle"] for a in appareils}

    # --- Emports (lus en jeu), par appareil et emplacement.
    emports = {}
    for r in lignes_emports:
        if r["json_key"] not in cles:
            continue
        sets = emports.setdefault(r["json_key"], {})
        idx = int(r["set_index"])
        s = sets.get(idx)
        if s is None:
            s = sets[idx] = {"index": idx, "nom_jeu": r["set_name"], "points": int(nombre(r["points"]) or 0),
                             "soute": r["soute"] == "1", "exclut_idx": [int(x) for x in r["exclut"].split("|") if x != ""],
                             "integres": int(nombre(r["integres"]) or 0), "options": []}
        if r["mount_jsonkey"]:
            cats = [c for c in r["categorie"].split("|") if c]
            s["options"].append({"cle": r["mount_jsonkey"], "nom": r["mount_name"], "masse": nombre(r["masse"]), "categories": cats})
    for a in appareils:
        sets = [emports.get(a["cle"], {})[i] for i in sorted(emports.get(a["cle"], {}))]
        noms = {s["index"]: s["nom_jeu"] for s in sets}
        for s in sets:
            s["exclut"] = [noms.get(i, f"#{i}") for i in s["exclut_idx"]]
            armees = [o for o in s["options"] if o["categories"]]
            pesees = [o["masse"] for o in armees if "CARGO" not in o["categories"] and o["masse"]]
            s["charge_kg"] = max(pesees) if pesees else None                       # A9 amendé v3.2
            s["categories"] = [c for c in CATEGORIES if any(c in o["categories"] for o in armees)]
            s["nb_armes"] = len(armees)
            s["armes"] = [o["nom"] for o in s["options"]]
            s["affiche"] = bool(armees)                                            # A23
            s["masquer"] = False
            s["nom_fr"] = classeur["emports"].get((a["cle"], s["nom_jeu"]))
            s["force"] = {}
        a["emplacements"] = sets
        if a["canons_fixes"] is None:
            # Repli lu en jeu (A4) : points des emplacements internes dont toutes les options sont des canons.
            canons = sum(s["points"] for s in sets if "internal" in s["nom_jeu"].lower() and s["options"]
                         and all(o["categories"] == ["CANON"] for o in s["options"]))
            a["canons_fixes"] = float(canons)
            if canons:
                a["notes_auto"].append(f"Canons fixes = {canons} (emplacements internes à canon, lus en jeu).")
        ed = classeur["avions"].get(a["cle"], {})
        for champ in ("roles", "description_fr", "texte_perso", "tags"):
            a[champ] = ed.get(champ)

    # --- A15 : masse à vide manquante.
    paires = [(a["masse_vide_kg"], a["masse_parts"]) for a in appareils if a["masse_vide_kg"] and a["masse_parts"]]
    concordent = sum(1 for mv, mp in paires if abs(mp - mv) <= 0.15 * mv)
    a15_ok = bool(paires) and concordent >= 0.8 * len(paires)
    for a in appareils:
        if not a["masse_vide_kg"]:
            if a15_ok and a["masse_parts"]:
                a["masse_vide_kg"] = a["masse_parts"]
                a["notes_auto"].append(f"Masse à vide = somme des pièces du prefab (A15 : {concordent}/{len(paires)} appareils concordants à ±15 %).")
            else:
                a["masse_vide_kg"] = None
                a["notes_auto"].append(f"Masse à vide inconnue : N/A (A15 : somme des pièces concordante pour {concordent}/{len(paires)} "
                                       f"appareils seulement, moins de 80 %).")

    # --- A20 : corrections, appliquées avant tout calcul.
    etats = appliquer_corrections(appareils, classeur["corrections"], journal)

    # --- Calculs (sur les valeurs corrigées).
    for a in appareils:
        if "charge_utile_kg" not in a["corriges"]:
            a["charge_utile_kg"] = charge_utile_emports(a["emplacements"])          # A14 amendé v3.6
        a["encombrement_m2"] = arrondi(a["longueur_m"] * a["envergure_m"]) if a["longueur_m"] and a["envergure_m"] else None
        a["ir_30"] = (arrondi(a["ir_min"] + 0.3 * (a["ir_max_sec"] - a["ir_min"]), 2)
                      if a["ir_min"] is not None and a["ir_max_sec"] is not None else None)
        a["note_rcs"] = note_rcs(a["rcs_m2"])
    seuils = {"encombrement_m2": tiers_log(a["encombrement_m2"] for a in appareils),
              "charge_kg": tiers_log(s["charge_kg"] for a in appareils for s in a["emplacements"] if s["affiche"] and not s["masquer"])}
    for a in appareils:
        if "taille" not in a["corriges"]:
            a["taille"] = classe(a["encombrement_m2"], seuils["encombrement_m2"])
        for s in a["emplacements"]:
            s["classe"] = classe(s["charge_kg"], seuils["charge_kg"])
        visibles = [s for s in a["emplacements"] if s["affiche"] and not s["masquer"]]
        a["nb_emplacements"] = len(visibles)
        a["nb_points"] = sum(s["points"] for s in visibles)
        a["nb_soutes"] = sum(1 for s in visibles if s["soute"])
        a["charge_max_point_kg"] = max((s["charge_kg"] for s in visibles if s["charge_kg"]), default=None)
    for champ, note in (("vitesse_kmh", "note_vitesse"), ("charge_utile_kg", "note_charge"), ("cout_m", "note_cout")):
        valeurs = [a[champ] for a in appareils]
        for a in appareils:
            a[note] = note_rang(a[champ], valeurs)
    appareils.sort(key=lambda a: (a["rang"] if a["rang"] is not None else 99, a["cout_m"] if a["cout_m"] is not None else 1e9, a["cle"]))
    return {"appareils": appareils, "seuils": seuils, "corrections": etats,
            "exports": (os.path.basename(export_avions), os.path.basename(export_emports))}


def _appareil(e, n):
    """Fusion champ par champ : NO-stats (Excel de Florian) d'abord, export du jeu en repli."""
    n = n or {}

    def nb(cle_no, cle_exp=None, zero_absent=False):
        v = nombre(n.get(cle_no)) if cle_no else None
        if v is None and cle_exp:
            v = nombre(e.get(cle_exp))
        if zero_absent and v == 0:
            v = None
        return v

    longueur, envergure, hauteur = dimensions(n.get("dimensions"))
    vertical = booleen(n.get("decollage_vertical"))
    if vertical is None:
        vertical = booleen(e.get("vertical"))
    notes = texte(n.get("notes"))
    return {
        "cle": e["json_key"], "nom": texte(n.get("nom")) or e["unit_name"], "code": texte(n.get("code")) or e["code"],
        "source": texte(n.get("source")) or "Mod (hors NO-stats)",
        "type": texte(n.get("type")) or type_depuis_moteurs(e.get("moteurs")),
        "rang": nb("rang", "rang"), "cout_m": nb("cout_m", "cout"), "vitesse_kmh": nb("vitesse_kmh", "vmax_info", True),
        "decrochage_kmh": nb("decrochage_kmh", "decrochage", True), "maniabilite": nb("maniabilite", "maniabilite"),
        "g_max": nb("g_max", "g_max"), "masse_vide_kg": nb("masse_vide_kg", "masse_vide", True),
        "masse_max_kg": nb("masse_max_kg", "masse_max", True), "masse_parts": nombre(e.get("masse_parts")),
        "carburant_kg": nb("carburant_kg"), "poussee_kn": nb("poussee_kn"), "puissance_kw": nb("puissance_kw"),
        "decollage_vertical": vertical, "decollage_m": nb("decollage_m", "decollage_m", True),
        "decollage_ms": nb("decollage_ms", "decollage_ms", True), "radar_m": nb("radar_m"),
        "brouilleur": booleen(n.get("brouilleur")), "canons_fixes": nb("canons_fixes"), "leurres": nb("leurres"),
        "paillettes": nb("paillettes"), "blindage": nb("blindage", "blindage"), "rcs_m2": nb("rcs_m2", "rcs", True),
        "longueur_m": longueur or nombre(e.get("longueur")), "envergure_m": envergure or nombre(e.get("largeur")),
        "hauteur_m": hauteur or nombre(e.get("hauteur")),
        "laser_cibles": nombre(e.get("laser_cibles")), "laser_portee_m": nombre(e.get("laser_portee")),
        "ir_min": nombre(e.get("ir_min")), "ir_max_sec": nombre(e.get("ir_max_sec")), "ir_max": nombre(e.get("ir_max")),
        "ir_force": {"ir_30": None, "ir_max": None},
        "role_ia": {k: nb(k, k) for k in ("role_surface", "role_air", "role_missile", "role_radar")},
        "description_jeu": texte(n.get("description_jeu")) or texte(e.get("description")),
        "notes": notes, "notes_auto": [], "corriges": [], "sources_corrigees": {},
    }


def charge_utile_emports(emplacements):
    """
    Amendement v3.6 (demande de Florian) : charge utile = masse d'armement maximale emportable sur l'ensemble des
    emplacements. Par emplacement : points × charge max par point (option d'arme la plus lourde, cargo exclu, A9) ;
    puis meilleure combinaison d'emplacements compatibles (exclusions mutuelles respectées). None si aucune arme pesée.
    """
    sets = [s for s in emplacements if not s.get("masquer")]
    poids = {s["index"]: (s["charge_kg"] or 0) * s["points"] for s in sets}
    excl = {s["index"]: set(s["exclut_idx"]) for s in sets}
    libres = [i for i in poids if not excl[i] and not any(i in excl[j] for j in poids)]
    lies = [i for i in poids if i not in libres]
    meilleur = 0.0
    for masque in range(1 << len(lies)):                        # peu d'emplacements liés : force brute
        choix = [lies[k] for k in range(len(lies)) if masque >> k & 1]
        if any(j in excl[i] or i in excl[j] for i in choix for j in choix if i != j):
            continue
        meilleur = max(meilleur, sum(poids[i] for i in choix))
    total = meilleur + sum(poids[i] for i in libres)
    return total or None


# ---------------------------------------------------------------------------------------------------------------
def convertir(valeur, typ):
    """Valeur forcée → valeur typée, ou (None, raison)."""
    if typ == "N":
        v = nombre(valeur)
        return (v, None) if v is not None else (None, "valeur non numérique")
    if typ == "B":
        v = booleen(valeur)
        return (v, None) if v is not None else (None, "valeur attendue : OUI ou NON")
    if typ == "L":
        long = {v: k for k, v in ACRONYMES.items()}
        cats = [long.get(c.strip().upper(), c.strip().upper()) for c in str(valeur or "").split(",") if c.strip()]
        inconnues = [c for c in cats if c not in CATEGORIES]
        if inconnues:
            return None, "catégorie inconnue : " + ", ".join(inconnues)
        return [c for c in CATEGORIES if c in cats], None
    v = texte(valeur)
    return (v, None) if v is not None else (None, "valeur vide")


def appliquer_corrections(appareils, lignes, journal=print):
    """A20 : applique chaque ligne valide ; renvoie l'état de chaque ligne (« appliquée » / « ignorée : … »)."""
    par_cle = {a["cle"]: a for a in appareils}
    etats, derniere = [], {}
    for i, (cle, emplacement, champ_lib, valeur, commentaire) in enumerate(lignes):
        cle, emplacement, champ_lib = texte(cle), texte(emplacement), texte(champ_lib)
        etat = None
        a = par_cle.get(cle)
        if a is None:
            etat = "ignorée : clé inconnue"
        elif champ_lib not in CHAMPS:
            etat = "ignorée : champ hors liste"
        else:
            champ, typ, de_emplacement = CHAMPS[champ_lib]
            if de_emplacement and not emplacement:
                etat = "ignorée : emplacement manquant pour ce champ"
            elif not de_emplacement and emplacement:
                etat = "ignorée : ce champ concerne l'appareil, laisse « Emplacement » vide"
            else:
                v, raison = convertir(valeur, typ)
                if raison:
                    etat = "ignorée : " + raison
                elif champ == "taille" and v.upper() not in ("S", "M", "H"):
                    etat = "ignorée : taille attendue S, M ou H"
                elif de_emplacement:
                    s = next((s for s in a["emplacements"] if s["nom_jeu"] == emplacement), None)
                    if s is None:
                        etat = "ignorée : emplacement inconnu"
                    else:
                        cible = (cle, emplacement, champ)
                        if cible in derniere:
                            etats[derniere[cible]] = f"ignorée : remplacée par la ligne {i + 2}"
                        derniere[cible] = i
                        s["force"][champ] = v
                        if champ == "masquer":
                            s["masquer"] = v
                        else:
                            s[champ] = v
                        etat = "appliquée"
                else:
                    cible = (cle, None, champ)
                    if cible in derniere:
                        etats[derniere[cible]] = f"ignorée : remplacée par la ligne {i + 2}"
                    derniere[cible] = i
                    if champ == "taille":
                        v = v.upper()
                    if champ.startswith("ir_force."):
                        a["ir_force"][champ.split(".", 1)[1]] = v
                    else:
                        a["sources_corrigees"].setdefault(champ, a.get(champ))
                        a[champ] = v
                    if champ not in a["corriges"]:
                        a["corriges"].append(champ)
                    etat = "appliquée"
        etats.append(etat)
        if etat != "appliquée":
            journal(f"Avertissement : correction ligne {i + 2} ({cle} · {emplacement or '—'} · {champ_lib}) {etat}.")
    # Une correction remplacée plus loin a déjà été écrasée par la suivante : rien à défaire.
    return etats
